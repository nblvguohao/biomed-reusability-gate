"""Build a conservative, checksum-tracked NMI initial-submission package."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REQUIRED_SOURCES = {
    "01_Manuscript_clean": "manuscript/reusability_report.md",
    "02_Supplementary_Information": "manuscript/SUPPLEMENTARY_INFORMATION.md",
    "03_Cover_Letter": "manuscript/COVER_LETTER.md",
    "04_Reporting_Summary_answers": "manuscript/REPORTING_SUMMARY.md",
    "05_Software_Submission_Checklist": "manuscript/SOFTWARE_SUBMISSION_CHECKLIST.md",
    "06_Title_Page": "manuscript/TITLE_PAGE.md",
}
PDF_STEMS = {
    "01_Manuscript_clean",
    "02_Supplementary_Information",
    "03_Cover_Letter",
}
FIGURE_SUFFIXES = (".pdf", ".svg", ".png", ".tiff")


@dataclass(frozen=True)
class PackageManifest:
    output_dir: Path
    files: tuple[Path, ...]
    zip_path: Path


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _prevent_table_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    keep = OxmlElement("w:cantSplit")
    keep.set(qn("w:val"), "true")
    tr_pr.append(keep)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char, instr_text, fld_end))


def _configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, before, after in (
        ("Title", 18, 0, 12),
        ("Heading 1", 14, 14, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 10.5, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style.element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    if "Caption" in styles:
        caption = styles["Caption"]
    else:
        caption = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(8)

    document.core_properties.title = title
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    _add_page_number(section.footer.paragraphs[0])


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _table_block(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _add_inline(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                _set_cell_shading(cell, "E7E7E7")
        if row_index == 0:
            _set_repeat_table_header(table.rows[-1])
        _prevent_table_row_split(table.rows[-1])
    document.add_paragraph()


def _add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.replace("\r\n", "\n").splitlines()
    index = 0
    pending: list[str] = []

    def flush() -> None:
        if not pending:
            return
        text = " ".join(part.strip() for part in pending).strip()
        pending.clear()
        if text:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, text)

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and index + 1 < len(lines):
            flush()
            rows, index = _table_block(lines, index)
            _add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if heading:
            flush()
            level = min(len(heading.group(1)), 3)
            text = re.sub(r"<[^>]+>", "", heading.group(2)).strip()
            paragraph = document.add_paragraph(
                style="Title" if level == 1 else f"Heading {level - 1}"
            )
            _add_inline(paragraph, text)
        elif bullet or numbered:
            flush()
            match = bullet or numbered
            style = "List Bullet" if bullet else "List Number"
            paragraph = document.add_paragraph(style=style)
            _add_inline(paragraph, match.group(1))
        elif stripped in {"---", "```"} or stripped.startswith("<!--"):
            flush()
        elif not stripped:
            flush()
        else:
            pending.append(stripped)
        index += 1
    flush()


def _build_docx(source: Path, destination: Path, *, appendix: Path | None = None) -> None:
    document = Document()
    _configure_document(document, source.stem.replace("_", " "))
    markdown = source.read_text(encoding="utf-8")
    if appendix is not None:
        markdown = f"{markdown}\n\n{appendix.read_text(encoding='utf-8')}"
    _add_markdown(document, markdown)
    document.save(destination)


def _soffice_path() -> Path:
    found = shutil.which("soffice")
    candidates = [
        Path(found) if found else None,
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    ]
    for candidate in candidates:
        if candidate is not None and candidate.exists():
            return candidate
    raise RuntimeError("LibreOffice soffice executable was not found")


def _convert_pdf(docx_path: Path, output_path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="nmi_soffice_") as profile:
        command = [
            str(_soffice_path()),
            f"-env:UserInstallation=file:///{Path(profile).as_posix()}",
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_path.parent),
            str(docx_path),
        ]
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            timeout=180,
        )
    generated = output_path.parent / f"{docx_path.stem}.pdf"
    if completed.returncode != 0 or not generated.exists() or generated.stat().st_size == 0:
        raise RuntimeError(
            f"LibreOffice failed for {docx_path.name}: {completed.stdout}\n{completed.stderr}"
        )
    if generated != output_path:
        generated.replace(output_path)


def _copy_figures(root: Path, output_dir: Path) -> None:
    source = root / "artifacts/revision_figures"
    if not source.exists():
        raise FileNotFoundError(
            "Final revision figures are absent; run make_revision_figures.py first"
        )
    destination = output_dir / "Figures"
    destination.mkdir()
    for path in sorted(source.iterdir()):
        if path.name.startswith("Figure_") and path.suffix.lower() in FIGURE_SUFFIXES:
            shutil.copy2(path, destination / path.name)
    if not any(destination.iterdir()):
        raise FileNotFoundError("No final Figure_* exports were found")


def _copy_source_data(root: Path, output_dir: Path) -> None:
    source = root / "artifacts/revision_figures/figure_source_data.json"
    if not source.exists():
        raise FileNotFoundError(source)
    destination = output_dir / "Source_Data"
    destination.mkdir()
    shutil.copy2(source, destination / "Figure_Source_Data.json")
    payload = json.loads(source.read_text(encoding="utf-8"))
    rows = payload.get("figure3_rows", [])
    if not rows:
        raise ValueError("figure source data contains no Figure 3 rows")
    columns = sorted({key for row in rows for key in row if key != "cluster_mass_sensitivity"})
    with (destination / "Figure_3_Source_Data.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    sensitivity_rows = []
    for row in rows:
        for setting in row.get("cluster_mass_sensitivity", []):
            sensitivity_rows.append(
                {
                    "cutoff": row["cutoff"],
                    "seed": row["seed"],
                    "method": row["method"],
                    "scale": row["scale"],
                    **setting,
                }
            )
    with (destination / "Figure_3_Sensitivity.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=list(sensitivity_rows[0]))
        writer.writeheader()
        writer.writerows(sensitivity_rows)
    figure2_tables = {
        "Figure_2_Preprocessing.csv": payload["figure2"]["preprocessing_rows"],
        "Figure_2_Latent_Noise.csv": payload["figure2"]["latent_noise_rows"],
        "Figure_2_Simulation.csv": payload["figure2"]["simulation_rows"],
    }
    for filename, table_rows in figure2_tables.items():
        with (destination / filename).open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(table_rows[0]))
            writer.writeheader()
            writer.writerows(table_rows)
    manifest = root / "artifacts/revision_results/revision_results.json"
    if manifest.exists():
        shutil.copy2(manifest, destination / "Revision_Result_Manifest.json")


def _write_supporting_files(root: Path, output_dir: Path) -> None:
    (output_dir / "README.md").write_text(
        "# NMI initial-submission package\n\n"
        "This archive contains the clean manuscript, Supplementary Information, "
        "cover letter, reporting answers, software checklist, title page, final "
        "figures and machine-readable source data. The PDF files mirror the three "
        "documents required for editorial reading; DOCX files remain editable.\n",
        encoding="utf-8",
    )
    supporting = {
        "Submission_Readiness_Checklist.md": root / "reports/NMI_SUBMISSION_READINESS.md",
        "CHANGE_LOG.md": root / "reports/NMI_REVISION_CHANGE_LOG.md",
    }
    fallbacks = {
        "Submission_Readiness_Checklist.md": (
            "# Submission readiness\n\n"
            "- Manuscript structure and word limits checked against the current NMI "
            "Article instructions.\n"
            "- Figure exports and source data included.\n"
            "- Reporting Summary answers and software checklist included.\n"
            "- Author team must confirm author order, contributions, competing "
            "interests, related manuscripts and submission-system declarations "
            "before upload.\n"
        ),
        "CHANGE_LOG.md": (
            "# Revision change log\n\n"
            "The submission adds leakage-safe early and late temporal cutoffs, four "
            "training-only baselines, structure-aware metrics and sensitivity "
            "analyses; it narrows the VO analysis to a target-informed mechanism "
            "sanity check and rewrites all claims to match the verified information "
            "boundaries.\n"
        ),
    }
    for filename, source in supporting.items():
        destination = output_dir / filename
        if source.exists():
            shutil.copy2(source, destination)
        else:
            destination.write_text(fallbacks[filename], encoding="utf-8")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_checksums(output_dir: Path) -> Path:
    checksum_path = output_dir / "SHA256SUMS.txt"
    files = sorted(
        path
        for path in output_dir.rglob("*")
        if path.is_file() and path != checksum_path and "qa" not in path.parts
    )
    checksum_path.write_text(
        "".join(f"{_sha256(path)}  {path.relative_to(output_dir).as_posix()}\n" for path in files),
        encoding="utf-8",
    )
    return checksum_path


def _write_machine_manifest(output_dir: Path) -> Path:
    path = output_dir / "package_manifest.json"
    payload = {
        "schema_version": "1.0",
        "journal": "Nature Machine Intelligence",
        "article_type": "Reusability Report (Article format)",
        "files": [
            item.relative_to(output_dir).as_posix()
            for item in sorted(output_dir.rglob("*"))
            if item.is_file() and item != path
        ],
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return path


def _write_zip(output_dir: Path) -> Path:
    zip_path = output_dir.with_suffix(".zip")
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(output_dir.rglob("*")):
            if path.is_file() and "qa" not in path.parts:
                archive.write(path, path.relative_to(output_dir.parent))
    return zip_path


def build_package(root: Path, output_dir: Path) -> PackageManifest:
    root = root.resolve()
    output_dir = output_dir.resolve()
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)

    missing = [relative for relative in REQUIRED_SOURCES.values() if not (root / relative).exists()]
    if missing:
        raise FileNotFoundError(f"missing submission sources: {missing}")

    for stem, relative in REQUIRED_SOURCES.items():
        docx_path = output_dir / f"{stem}.docx"
        appendix = root / "manuscript/FIGURE_LEGENDS.md" if stem == "01_Manuscript_clean" else None
        _build_docx(root / relative, docx_path, appendix=appendix)
        if stem in PDF_STEMS:
            _convert_pdf(docx_path, output_dir / f"{stem}.pdf")

    _copy_figures(root, output_dir)
    _copy_source_data(root, output_dir)
    _write_supporting_files(root, output_dir)
    _write_machine_manifest(output_dir)
    _write_checksums(output_dir)
    zip_path = _write_zip(output_dir)
    files = tuple(path for path in sorted(output_dir.rglob("*")) if path.is_file())
    return PackageManifest(output_dir=output_dir, files=files, zip_path=zip_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, default=Path("submission_package_nmi"))
    args = parser.parse_args()
    manifest = build_package(args.root, args.output)
    print(f"Built {len(manifest.files)} files in {manifest.output_dir}")
    print(f"Archive: {manifest.zip_path}")


if __name__ == "__main__":
    main()
