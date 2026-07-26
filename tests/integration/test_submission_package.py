"""Integration contract for the NMI initial-submission package."""

from __future__ import annotations

import hashlib
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from build_submission_package import build_package  # noqa: E402


@pytest.fixture(scope="module")
def package_manifest(tmp_path_factory):
    return build_package(ROOT, tmp_path_factory.mktemp("package") / "submission")


def test_package_contains_required_deliverables(package_manifest):
    manifest = package_manifest
    names = {path.relative_to(manifest.output_dir).as_posix() for path in manifest.files}

    assert {
        "01_Manuscript_clean.docx",
        "01_Manuscript_clean.pdf",
        "02_Supplementary_Information.docx",
        "02_Supplementary_Information.pdf",
        "03_Cover_Letter.docx",
        "03_Cover_Letter.pdf",
        "04_Reporting_Summary_answers.docx",
        "05_Software_Submission_Checklist.docx",
        "06_Title_Page.docx",
        "README.md",
        "Submission_Readiness_Checklist.md",
        "CHANGE_LOG.md",
        "SHA256SUMS.txt",
        "Source_Data/Figure_3_Source_Data.csv",
        "Source_Data/Figure_3_Sensitivity.csv",
        "Source_Data/Figure_2_Preprocessing.csv",
        "Source_Data/Figure_2_Latent_Noise.csv",
    }.issubset(names)
    manuscript = Document(manifest.output_dir / "01_Manuscript_clean.docx")
    manuscript_text = "\n".join(paragraph.text for paragraph in manuscript.paragraphs)
    assert "Fig. 1 |" in manuscript_text


def test_docx_tables_repeat_headers_and_do_not_split_rows(package_manifest):
    for docx_path in package_manifest.output_dir.glob("*.docx"):
        document = Document(docx_path)
        for table in document.tables:
            assert table.rows
            header_properties = table.rows[0]._tr.get_or_add_trPr()
            assert header_properties.find(qn("w:tblHeader")) is not None
            for row in table.rows:
                row_properties = row._tr.get_or_add_trPr()
                assert row_properties.find(qn("w:cantSplit")) is not None


def test_package_excludes_archival_sources_and_private_paths(package_manifest):
    manifest = package_manifest
    package_text = "\n".join(
        path.read_text(encoding="utf-8", errors="ignore")
        for path in manifest.files
        if path.suffix.lower() in {".md", ".txt", ".json", ".svg"}
    ).lower()
    names = "\n".join(path.name.lower() for path in manifest.files)

    assert "reusability_report_draft" not in names
    assert "lab_host" not in package_text
    assert "c:\\cc\\" not in package_text
    assert "/data/lgh/" not in package_text


def test_package_checksums_match_files(package_manifest):
    manifest = package_manifest
    listed = {}
    for line in (manifest.output_dir / "SHA256SUMS.txt").read_text(encoding="utf-8").splitlines():
        digest, relative = line.split("  ", 1)
        listed[relative] = digest

    for relative, expected in listed.items():
        path = manifest.output_dir / relative
        assert path.exists()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == expected

    with zipfile.ZipFile(manifest.zip_path) as archive:
        names = set(archive.namelist())
        prefix = f"{manifest.output_dir.name}/"
        assert f"{prefix}01_Manuscript_clean.docx" in names
        assert all("reusability_report_draft" not in name.lower() for name in names)
